import json
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]


def test_executed_pipeline_outputs():
    stats = json.loads((ROOT / "reports/evaluation.json").read_text())
    assert stats["raw_records"] == stats["accepted_records"] + stats["rejected_records"]
    assert stats["silver_units"] == stats["gold_units"]
    gold = ROOT / "data/processed/gold"
    daily = pd.read_csv(gold / "daily_demand.csv")
    assert {"product_id", "date", "units"} <= set(daily.columns)
    assert len(daily) == stats["products"] * stats["history_days"]
    assert not daily.duplicated(["product_id", "date"]).any()
    for name in [
        "product_master",
        "daily_demand",
        "weekly_demand",
        "date_dimension",
        "forecast_results",
        "future_forecast",
        "inventory_risk",
        "product_forecast_scores",
    ]:
        df = pd.read_csv(gold / f"{name}.csv")
        assert len(df)
        required = {
            "product_master": {"product_id", "description"},
            "date_dimension": {"date", "year", "month", "week"},
            "forecast_results": {
                "product_id",
                "date",
                "train_end",
                "model",
                "actual",
                "predicted",
                "absolute_error",
            },
            "future_forecast": {"product_id", "date", "as_of", "model", "predicted"},
            "inventory_risk": {
                "product_id",
                "as_of",
                "stock_risk",
                "days_cover",
                "planning_source",
            },
        }.get(name, {"product_id"})
        assert required <= set(df.columns)
        assert len(
            pd.read_parquet(gold / name if (gold / name).is_dir() else gold / f"{name}.parquet")
        ) == len(df)
    checks = pd.read_csv(ROOT / "reports/kpi_reconciliation.csv")
    assert checks.passed.all()
    assert checks.difference.abs().max() < 1e-6


def test_excel_structure_and_cached_formulas():
    import openpyxl

    path = ROOT / "excel/SupplySight_Operations_Analysis.xlsx"
    assert path.exists(), "Generate workbook before full release checks"
    wb = openpyxl.load_workbook(path, data_only=False)
    expected = {
        "Executive Summary",
        "Product Master",
        "Demand History",
        "Forecast Results",
        "Inventory Risk",
        "KPI Validation",
        "Data Dictionary",
    }
    expected.update({"Capacity Planning", "Forecast Study"})
    assert set(wb.sheetnames) == expected
    assert len(wb["Executive Summary"]._charts) == 1
    assert len(wb["Inventory Risk"].conditional_formatting) > 0
    calculated = openpyxl.load_workbook(path, data_only=True)
    stats = json.loads((ROOT / "reports/evaluation.json").read_text())
    assert calculated["Executive Summary"]["B6"].value == stats["gold_units"]
    assert calculated["Executive Summary"]["B7"].value == stats["products"]
    assert calculated["Executive Summary"]["B8"].value == stats["at_risk_products_synthetic"]
    assert calculated["Executive Summary"]["B9"].value == stats["overstock_products_synthetic"]
    model_scores = {x["model"]: x for x in stats["model_comparison"]}
    for row in range(14, 17):
        name = calculated["Executive Summary"].cell(row, 1).value
        assert calculated["Executive Summary"].cell(row, 3).value == pytest.approx(
            model_scores[name]["wape"]
        )
    for row in calculated["KPI Validation"].iter_rows(
        min_row=6, max_row=5 + stats["reconciled_kpis"]
    ):
        assert abs(row[3].value) < 1e-6
    for sheet in calculated:
        for row in sheet:
            for cell in row:
                assert cell.data_type != "e", f"{sheet.title}!{cell.coordinate}: {cell.value}"
